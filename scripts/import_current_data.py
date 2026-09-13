from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import sys

from openpyxl import load_workbook
from sqlalchemy import select

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE))

from app.database import SessionLocal
from app.models import (
    Branch,
    DailyMenu,
    Ingredient,
    Recipe,
    Supplier,
    Seasoning,
)


CURRENT = BASE / "data" / "current"


def clean(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date()
    return str(value).strip()


def parse_date(value):
    value = clean(value)
    if isinstance(value, date):
        return value
    if not value:
        return None
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Unsupported date format: {value}")


def rows(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    header = [clean(cell.value) for cell in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {header[index]: clean(value) for index, value in enumerate(row) if index < len(header)}
        if any(item.values()):
            yield item


def first_data_row(path: Path):
    return next(rows(path))


def get_branch(db):
    menu_file = CURRENT / "menu_20260912115546.xlsx"
    sample = first_data_row(menu_file)
    school_name = sample["學校名稱*"]
    service_location = sample["供餐地點*"]
    restaurant_name = sample["餐廳名稱*"]

    branch = db.scalar(
        select(Branch).where(
            Branch.school_name == school_name,
            Branch.service_location == service_location,
            Branch.restaurant_name == restaurant_name,
        )
    )
    if not branch:
        branch = db.scalar(select(Branch).where(Branch.id == 1))
    if branch:
        branch.name = restaurant_name
        branch.school_name = school_name
        branch.service_location = service_location
        branch.restaurant_name = restaurant_name
        branch.active = True
    else:
        branch = Branch(
            name=restaurant_name,
            school_name=school_name,
            service_location=service_location,
            restaurant_name=restaurant_name,
            active=True,
        )
        db.add(branch)
    db.flush()
    return branch


def get_or_create_supplier(db, name, branch_id=None, **attrs):
    name = clean(name)
    if not name:
        return None
    supplier = db.scalar(select(Supplier).where(Supplier.name == name))
    if not supplier:
        supplier = Supplier(name=name, branch_id=branch_id)
        db.add(supplier)
    for key, value in attrs.items():
        if value and not getattr(supplier, key):
            setattr(supplier, key, value)
    db.flush()
    return supplier


def import_suppliers(db, branch):
    count = 0
    for item in rows(CURRENT / "供應商資料.xlsx"):
        supplier = get_or_create_supplier(
            db,
            item["供應商名稱"],
            branch.id,
            owner=item.get("負責人") or None,
            tax_id=item.get("公司統編") or None,
            address=item.get("地址") or None,
            phone=item.get("電話") or None,
        )
        if supplier:
            count += 1
    return count


def import_ingredients(db, branch):
    count = 0
    path = CURRENT / "schoolrestaurantingredient_20260912115642.xlsx"
    for item in rows(path):
        supplier = get_or_create_supplier(db, item["供應商名稱*"], branch.id)
        if not supplier:
            continue
        product_name = item["產品名稱*"]
        ingredient_name = item["食材名稱*"]
        origin = item["食材原產地(國)*"]
        ingredient = db.scalar(
            select(Ingredient).where(
                Ingredient.product_name == product_name,
                Ingredient.ingredient_name == ingredient_name,
                Ingredient.supplier_id == supplier.id,
                Ingredient.branch_id == branch.id,
            )
        )
        if not ingredient:
            ingredient = Ingredient(
                branch_id=branch.id,
                product_name=product_name,
                ingredient_name=ingredient_name,
                origin=origin,
                supplier_id=supplier.id,
            )
            db.add(ingredient)
            count += 1
        ingredient.manufacturer = item.get("製造商") or None
        ingredient.processed_food = item.get("加工食品") or None
        ingredient.non_gmo_soy = item.get("非基改黃豆") or None
        ingredient.non_gmo_corn = item.get("非基改玉米") or None
        ingredient.certification = item.get("食材驗證標章") or None
        ingredient.certification_no = item.get("驗證號碼") or None
        ingredient.unit = item.get("單位") or None
        ingredient.calories = item.get("熱量") or None
    return count


def import_seasonings(db, branch):
    count = 0
    path = CURRENT / "調味料new-輔大.xlsx"
    for item in rows(path):
        supplier = get_or_create_supplier(db, item["供應商*"], branch.id)
        if not supplier:
            continue
        name = item["調味料名稱*"]
        seasoning = db.scalar(
            select(Seasoning).where(
                Seasoning.name == name,
                Seasoning.supplier_id == supplier.id,
                Seasoning.branch_id == branch.id,
            )
        )
        if not seasoning:
            seasoning = Seasoning(branch_id=branch.id, name=name, supplier_id=supplier.id)
            db.add(seasoning)
            count += 1
        seasoning.manufacturer = item.get("製造商") or None
        seasoning.product_name = item.get("產品名稱") or None
        seasoning.certification = item.get("驗證標章") or None
        seasoning.certification_no = item.get("驗證號碼") or None
        seasoning.unit = item.get("單位") or None
        seasoning.non_gmo_soy = item.get("基改黃豆") or None
        seasoning.non_gmo_corn = item.get("基改玉米") or None
        seasoning.processed = item.get("加工品") or None
    return count


def import_menu(db, branch):
    count = 0
    path = CURRENT / "menu_20260912115546.xlsx"
    for item in rows(path):
        name = item["菜色名稱*"]
        recipe = db.scalar(select(Recipe).where(Recipe.name == name, Recipe.branch_id == branch.id))
        if not recipe:
            recipe = Recipe(
                branch_id=branch.id,
                name=name,
                category=item.get("菜色類別") or None,
                calories=int(float(item["熱量*"] or 0)),
            )
            db.add(recipe)
            db.flush()
        else:
            recipe.category = item.get("菜色類別") or None
            recipe.calories = int(float(item["熱量*"] or 0))

        service_date = parse_date(item["供餐日期*"])
        menu = db.scalar(
            select(DailyMenu).where(
                DailyMenu.branch_id == branch.id,
                DailyMenu.service_date == service_date,
                DailyMenu.recipe_id == recipe.id,
            )
        )
        if not menu:
            menu = DailyMenu(
                branch_id=branch.id,
                service_date=service_date,
                recipe_id=recipe.id,
                upload_state="pending",
            )
            db.add(menu)
            count += 1
    return count


def main():
    db = SessionLocal()
    try:
        branch = get_branch(db)
        result = {
            "branch": branch.restaurant_name,
            "suppliers_seen": import_suppliers(db, branch),
            "ingredients_added": import_ingredients(db, branch),
            "seasonings_added": import_seasonings(db, branch),
            "menu_items_added": import_menu(db, branch),
        }
        db.commit()
        print(result)
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    main()
