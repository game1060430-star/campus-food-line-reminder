from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (
    Branch,
    DailyMenu,
    Ingredient,
    ItemBranchScope,
    Recipe,
    RecipeIngredient,
    RecipeSeasoning,
    Seasoning,
    Supplier,
)


ROOT = Path(__file__).resolve().parents[2]
TEMPLATES = ROOT / "data" / "templates"
EXPORTS = ROOT / "exports"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


@dataclass
class ExportResult:
    files: dict[str, Path]
    errors: list[str]
    warnings: list[str]
    row_counts: dict[str, int]


def parse_excluded_dates(raw: str) -> set[date]:
    result: set[date] = set()
    for part in raw.replace("，", ",").replace("\n", ",").split(","):
        text = part.strip()
        if not text:
            continue
        result.add(datetime.strptime(text, "%Y-%m-%d").date())
    return result


def service_dates(start: date, end: date, weekdays: set[int], excluded: set[date]) -> list[date]:
    days: list[date] = []
    current = start
    while current <= end:
        if current.weekday() in weekdays and current not in excluded:
            days.append(current)
        current += timedelta(days=1)
    return days


def clone_row_style(sheet, source_row: int, target_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        src = sheet.cell(source_row, col)
        dst = sheet.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)


def reset_data_rows(sheet, start_row: int = 2) -> None:
    if sheet.max_row >= start_row:
        sheet.delete_rows(start_row, sheet.max_row - start_row + 1)


def required_column_indexes(sheet) -> set[int]:
    required = {
        col
        for col in range(1, sheet.max_column + 1)
        if any(marker in str(sheet.cell(1, col).value or "") for marker in ("*", "＊"))
    }
    return required or set(range(1, sheet.max_column + 1))


def column_letter(col: int) -> str:
    letters = ""
    while col:
        col, remainder = divmod(col - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def excel_date_number(value: date) -> int:
    return (value - date(1899, 12, 30)).days


def cell_text(value) -> str:
    if isinstance(value, datetime):
        return value.date().strftime("%Y/%m/%d")
    if isinstance(value, date):
        return value.strftime("%Y/%m/%d")
    return str(value)


def shared_string_values(shared_xml: str) -> list[str]:
    root = ET.fromstring(shared_xml.encode("utf-8"))
    values: list[str] = []
    for si in root.findall(f"{{{MAIN_NS}}}si"):
        values.append("".join(node.text or "" for node in si.iter(f"{{{MAIN_NS}}}t")))
    return values


def escape_xml_text(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def append_shared_strings(shared_xml: str, values: list[str]) -> tuple[str, dict[str, int]]:
    existing = shared_string_values(shared_xml)
    indexes = {value: index for index, value in enumerate(existing)}
    appended = []
    for value in values:
        if value not in indexes:
            indexes[value] = len(existing)
            existing.append(value)
            preserve = ' xml:space="preserve"' if value != value.strip() else ""
            appended.append(f"<si><t{preserve}>{escape_xml_text(value)}</t></si>")
    if appended:
        shared_xml = shared_xml.replace("</sst>", "".join(appended) + "</sst>")
    total = str(len(existing))
    shared_xml = re.sub(r'\bcount="\d+"', f'count="{total}"', shared_xml, count=1)
    shared_xml = re.sub(r'\buniqueCount="\d+"', f'uniqueCount="{total}"', shared_xml, count=1)
    return shared_xml, indexes


def worksheet_path_for_sheet(workbook_xml: str, rels_xml: str, sheet_name: str) -> str:
    workbook = ET.fromstring(workbook_xml.encode("utf-8"))
    rels = ET.fromstring(rels_xml.encode("utf-8"))
    rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
    for sheet in workbook.findall(f".//{{{MAIN_NS}}}sheet"):
        if sheet.attrib.get("name") == sheet_name:
            rel_id = sheet.attrib[f"{{{REL_NS}}}id"]
            target = rel_map[rel_id].lstrip("/")
            return target if target.startswith("xl/") else f"xl/{target}"
    raise KeyError(f"Template is missing required sheet: {sheet_name}")


def row_one_xml(sheet_xml: str) -> str:
    match = re.search(r'<row\b[^>]*\br="1"[^>]*>.*?</row>', sheet_xml, re.DOTALL)
    if not match:
        raise ValueError("Template sheet is missing header row.")
    return match.group(0)


def data_row_styles(sheet_xml: str) -> dict[int, str]:
    match = re.search(r'<row\b[^>]*\br="2"[^>]*>(.*?)</row>', sheet_xml, re.DOTALL)
    if not match:
        return {}
    styles: dict[int, str] = {}
    for cell in re.finditer(r'<c\b([^>]*)>', match.group(1)):
        attrs = cell.group(1)
        ref = re.search(r'\br="([A-Z]+)2"', attrs)
        style = re.search(r'\bs="([^"]+)"', attrs)
        if ref and style:
            col = 0
            for char in ref.group(1):
                col = col * 26 + ord(char) - 64
            styles[col] = style.group(1)
    return styles


def cell_xml(row_index: int, col: int, value, shared_indexes: dict[str, int], styles: dict[int, str], date_cols: set[int]) -> str:
    if value is None or value == "":
        return ""
    ref = f"{column_letter(col)}{row_index}"
    style = f' s="{styles[col]}"' if col in styles else ""
    if col in date_cols and isinstance(value, (date, datetime)):
        day = value.date() if isinstance(value, datetime) else value
        return f'<c r="{ref}"{style}><v>{excel_date_number(day)}</v></c>'
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f'<c r="{ref}"{style}><v>{value}</v></c>'
    index = shared_indexes[cell_text(value)]
    return f'<c r="{ref}"{style} t="s"><v>{index}</v></c>'


def write_rows_to_sheet_xml(sheet_xml: str, rows: list[list], columns_to_write: set[int], max_col: int, shared_indexes: dict[str, int], date_cols: set[int] | None) -> str:
    header = row_one_xml(sheet_xml)
    styles = data_row_styles(sheet_xml)
    effective_date_cols = date_cols or set()
    data_rows = []
    for row_index, row in enumerate(rows, start=2):
        cells = []
        for col in range(1, max_col + 1):
            if col not in columns_to_write:
                continue
            value = row[col - 1] if col <= len(row) else ""
            cell = cell_xml(row_index, col, value, shared_indexes, styles, effective_date_cols)
            if cell:
                cells.append(cell)
        data_rows.append(f'<row r="{row_index}" spans="1:{max_col}">{"".join(cells)}</row>')
    last_row = max(1, len(rows) + 1)
    sheet_xml = re.sub(r'<dimension ref="[^"]*"', f'<dimension ref="A1:{column_letter(max_col)}{last_row}"', sheet_xml, count=1)
    return re.sub(r'(<sheetData>).*?(</sheetData>)', rf'\1{header}{"".join(data_rows)}\2', sheet_xml, count=1, flags=re.DOTALL)


def write_rows(path: Path, sheet_name: str, rows: list[list], date_cols: set[int] | None = None, extra_columns: set[int] | None = None) -> int:
    wb = load_workbook(path, read_only=True)
    sheet = wb[sheet_name]
    max_col = sheet.max_column
    columns_to_write = required_column_indexes(sheet) | (extra_columns or set())
    wb.close()

    string_values = [
        cell_text(row[col - 1])
        for row in rows
        for col in range(1, max_col + 1)
        if col in columns_to_write
        and col <= len(row)
        and row[col - 1] not in (None, "")
        and not (date_cols and col in date_cols and isinstance(row[col - 1], (date, datetime)))
        and not isinstance(row[col - 1], (int, float))
    ]

    tmp_path = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path, "r") as source:
        workbook_xml = source.read("xl/workbook.xml").decode("utf-8")
        rels_xml = source.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        sheet_path = worksheet_path_for_sheet(workbook_xml, rels_xml, sheet_name)
        shared_xml = source.read("xl/sharedStrings.xml").decode("utf-8")
        shared_xml, shared_indexes = append_shared_strings(shared_xml, string_values)
        sheet_xml = source.read(sheet_path).decode("utf-8")
        sheet_xml = write_rows_to_sheet_xml(sheet_xml, rows, columns_to_write, max_col, shared_indexes, date_cols)
        with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == sheet_path:
                    data = sheet_xml.encode("utf-8")
                elif item.filename == "xl/sharedStrings.xml":
                    data = shared_xml.encode("utf-8")
                target.writestr(item, data)
    tmp_path.replace(path)
    return len(rows)


def copy_template(name: str, out_dir: Path, out_name: str) -> Path:
    src = TEMPLATES / name
    dst = out_dir / out_name
    dst.write_bytes(src.read_bytes())
    return dst


def safe_filename(text: str) -> str:
    forbidden = '<>:"/\\|?*'
    cleaned = "".join("_" if ch in forbidden else ch for ch in text).strip()
    return cleaned or "分店"


def explicit_scope_ids(db: Session, item_type: str, item_id: int) -> set[int]:
    return set(db.scalars(select(ItemBranchScope.branch_id).where(ItemBranchScope.item_type == item_type, ItemBranchScope.item_id == item_id)).all())


def available_for_branch(db: Session, item, item_type: str, branch_id: int) -> bool:
    explicit = explicit_scope_ids(db, item_type, item.id)
    if explicit:
        return branch_id in explicit
    return item.branch_id in (None, branch_id)


def visible_items(db: Session, model, item_type: str, branch_id: int, order_column) -> list:
    items = db.scalars(select(model).where(model.active == True).order_by(order_column)).all()
    return [item for item in items if available_for_branch(db, item, item_type, branch_id)]


def missing_supplier_fields(supplier: Supplier) -> list[str]:
    checks = [
        ("負責人", supplier.owner),
        ("統編", supplier.tax_id),
        ("地址", supplier.address),
        ("電話", supplier.phone),
    ]
    return [label for label, value in checks if not str(value or "").strip()]


def missing_ingredient_fields(ingredient: Ingredient, supplier: Supplier | None) -> list[str]:
    checks = [
        ("產品名稱", ingredient.product_name),
        ("食材名稱", ingredient.ingredient_name),
        ("原產地", ingredient.origin),
        ("供應商", supplier.name if supplier else None),
    ]
    return [label for label, value in checks if not str(value or "").strip()]


def parse_delivery_weekdays(raw: str | None) -> set[int]:
    return {int(part) for part in (raw or "").split(",") if part.strip().isdigit()}


def previous_workday(day: date) -> date:
    current = day - timedelta(days=1)
    while current.weekday() >= 5:
        current -= timedelta(days=1)
    return current


def purchase_date_for(service_date: date, supplier: Supplier) -> date:
    delivery_days = parse_delivery_weekdays(supplier.delivery_weekdays)
    if not delivery_days:
        return previous_workday(service_date)
    current = service_date - timedelta(days=1)
    for _ in range(14):
        if current.weekday() in delivery_days:
            return current
        current -= timedelta(days=1)
    return previous_workday(service_date)


def selected_menu_rows(db: Session, branch_id: int, dates: list[date], recipe_ids: list[int]) -> list[tuple[date, Recipe, int | None]]:
    if recipe_ids:
        recipes = db.scalars(
            select(Recipe)
            .where(Recipe.id.in_(recipe_ids), Recipe.active == True)
            .order_by(Recipe.name)
        ).all()
        recipes = [recipe for recipe in recipes if available_for_branch(db, recipe, "recipe", branch_id)]
        return [(day, recipe, None) for day in dates for recipe in recipes]

    menus = db.scalars(
        select(DailyMenu)
        .where(DailyMenu.branch_id == branch_id, DailyMenu.service_date.in_(dates))
        .order_by(DailyMenu.service_date, DailyMenu.recipe_id)
    ).all()
    recipes = {r.id: r for r in visible_items(db, Recipe, "recipe", branch_id, Recipe.name)}
    return [(menu.service_date, recipes[menu.recipe_id], menu.servings) for menu in menus if menu.recipe_id in recipes]


def build_export(
    db: Session,
    branch_id: int,
    start: date,
    end: date,
    weekdays: set[int],
    excluded_dates: set[date],
    recipe_ids: list[int],
    file_types: set[str] | None = None,
) -> ExportResult:
    file_types = file_types or {"menu", "ingredients", "seasonings", "suppliers"}
    branch = db.get(Branch, branch_id)
    if not branch:
        return ExportResult({}, ["找不到分店。"], [], {})
    if end < start:
        return ExportResult({}, ["結束日期不能早於開始日期。"], [], {})

    dates = service_dates(start, end, weekdays, excluded_dates)
    needs_menu_data = bool(file_types & {"menu", "ingredients", "seasonings"})
    if needs_menu_data and not dates:
        return ExportResult({}, ["日期區間內沒有可登錄日期。"], [], {})

    menu_items = selected_menu_rows(db, branch_id, dates, recipe_ids) if needs_menu_data else []
    if needs_menu_data and not menu_items:
        return ExportResult({}, ["沒有符合條件的菜單。請先排菜單，或在匯出頁勾選要登錄的菜色。"], [], {})

    ingredients = {i.id: i for i in visible_items(db, Ingredient, "ingredient", branch_id, Ingredient.ingredient_name)}
    seasonings = {s.id: s for s in visible_items(db, Seasoning, "seasoning", branch_id, Seasoning.name)}
    suppliers = {s.id: s for s in visible_items(db, Supplier, "supplier", branch_id, Supplier.name)}
    recipe_ingredient_links = db.scalars(select(RecipeIngredient)).all()
    recipe_seasoning_links = db.scalars(select(RecipeSeasoning)).all()

    ingredients_by_recipe: dict[int, list[Ingredient]] = {}
    for link in recipe_ingredient_links:
        ingredient = ingredients.get(link.ingredient_id)
        if ingredient:
            ingredients_by_recipe.setdefault(link.recipe_id, []).append(ingredient)

    seasonings_by_recipe: dict[int, list[Seasoning]] = {}
    for link in recipe_seasoning_links:
        seasoning = seasonings.get(link.seasoning_id)
        if seasoning:
            seasonings_by_recipe.setdefault(link.recipe_id, []).append(seasoning)

    errors: list[str] = []
    warnings: list[str] = []
    used_supplier_ids: set[int] = set()
    used_seasoning_ids: set[int] = set()
    used_ingredient_keys: set[tuple[date, int]] = set()
    used_seasoning_keys: set[int] = set()

    menu_rows: list[list] = []
    ingredient_rows: list[list] = []
    seasoning_rows: list[list] = []

    for day, recipe, servings in menu_items:
        recipe_ingredients = ingredients_by_recipe.get(recipe.id, [])
        recipe_seasonings = seasonings_by_recipe.get(recipe.id, [])
        if not recipe_ingredients:
            warnings.append(f"{day:%Y-%m-%d}「{recipe.name}」沒有綁定這個分店可用的食材；菜色仍會產生，食材檔不會新增這道菜的食材列。")

        ingredient_text = "、".join(i.ingredient_name for i in recipe_ingredients)
        menu_rows.append([
            branch.school_name,
            branch.service_location,
            branch.restaurant_name,
            day,
            recipe.category or "",
            recipe.name,
            ingredient_text,
            recipe.calories,
        ])

        for ingredient in recipe_ingredients:
            supplier = suppliers.get(ingredient.supplier_id)
            missing_ingredient = missing_ingredient_fields(ingredient, supplier)
            if missing_ingredient and "ingredients" in file_types:
                errors.append(f"食材「{ingredient.ingredient_name}」缺少必填欄位：{'、'.join(missing_ingredient)}。")
                continue
            if supplier and "suppliers" in file_types:
                used_supplier_ids.add(supplier.id)
            key = (day, ingredient.id)
            if key in used_ingredient_keys:
                continue
            used_ingredient_keys.add(key)
            if "ingredients" in file_types:
                purchase_date = purchase_date_for(day, supplier)
                ingredient_rows.append([
                branch.school_name,
                branch.service_location,
                branch.restaurant_name,
                day,
                purchase_date,
                ingredient.product_name,
                ingredient.ingredient_name,
                ingredient.origin,
                supplier.name,
                ingredient.manufacturer or "",
                "",
                "",
                "",
                ingredient.processed_food or "",
                ingredient.non_gmo_soy or "",
                ingredient.non_gmo_corn or "",
                ingredient.certification or "",
                ingredient.certification_no or "",
                "",
                ingredient.unit or "",
                ingredient.calories or "",
                servings or "",
                ])

        for seasoning in recipe_seasonings:
            supplier = suppliers.get(seasoning.supplier_id)
            if not supplier and file_types & {"seasonings", "suppliers"}:
                errors.append(f"調味料「{seasoning.name}」沒有啟用中的供應商。")
                continue
            if supplier and "suppliers" in file_types:
                used_supplier_ids.add(supplier.id)
            used_seasoning_ids.add(seasoning.id)
            if seasoning.id in used_seasoning_keys:
                continue
            used_seasoning_keys.add(seasoning.id)
            if "seasonings" in file_types:
                seasoning_rows.append([
                branch.school_name,
                branch.service_location,
                branch.restaurant_name,
                seasoning.name,
                start,
                "",
                "",
                start,
                end,
                supplier.name,
                seasoning.manufacturer or "",
                seasoning.certification or "",
                seasoning.certification_no or "",
                "",
                seasoning.product_name or "",
                "",
                seasoning.unit or "",
                seasoning.non_gmo_soy or "",
                seasoning.non_gmo_corn or "",
                seasoning.processed or "",
                ])

    if errors:
        return ExportResult({}, errors, warnings, {})

    supplier_rows = []
    if "suppliers" in file_types and not used_supplier_ids:
        used_supplier_ids={supplier.id for supplier in suppliers.values()}
    for supplier_id in sorted(used_supplier_ids):
        supplier = suppliers[supplier_id]
        missing = missing_supplier_fields(supplier)
        if missing:
            errors.append(f"供應商「{supplier.name}」缺少必填欄位：{'、'.join(missing)}。")
            continue
        supplier_rows.append([
            supplier.name,
            supplier.owner or "",
            supplier.tax_id or "",
            supplier.address or "",
            supplier.phone or "",
        ])

    if errors:
        return ExportResult({}, errors, warnings, {})

    if "seasonings" in file_types and not seasoning_rows:
        warnings.append("這次選取的菜色沒有綁定調味料，所以調味料檔只有表頭。")
    if "ingredients" in file_types and not ingredient_rows:
        warnings.append("這次選取的菜色沒有綁定食材，所以食材檔只有表頭。")

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = EXPORTS / f"official_excel_{stamp}"
    out_dir.mkdir(parents=True, exist_ok=True)

    prefix = safe_filename(branch.name)
    counts = {}
    files = {}
    if "menu" in file_types:
        menu_path = copy_template("PreMenuExcelExample.xlsx", out_dir, f"{prefix}_菜單_{start:%Y%m%d}_{end:%Y%m%d}.xlsx")
        counts["menu"] = write_rows(menu_path, "菜色清單", menu_rows, {4}, {7})
        files["菜單"] = menu_path
    if "ingredients" in file_types:
        ingredient_path = copy_template("PrerestaurantingredientExcelExample.xlsx", out_dir, f"{prefix}_食材_{start:%Y%m%d}_{end:%Y%m%d}.xlsx")
        counts["ingredients"] = write_rows(ingredient_path, "每日進貨食材", ingredient_rows, {4, 5, 11, 12})
        files["食材"] = ingredient_path
    if "seasonings" in file_types:
        seasoning_path = copy_template("seasoningstockdataCollegeExcelExample.xlsx", out_dir, f"{prefix}_調味料_{start:%Y%m%d}_{end:%Y%m%d}.xlsx")
        counts["seasonings"] = write_rows(seasoning_path, "Sheet1", seasoning_rows, {5, 6, 7, 8, 9})
        files["調味料"] = seasoning_path
    if "suppliers" in file_types:
        supplier_path = copy_template("supplierExcelExample.xlsx", out_dir, f"{prefix}_供應商.xlsx")
        counts["suppliers"] = write_rows(supplier_path, "Data", supplier_rows)
        files["供應商"] = supplier_path
    return ExportResult(files, [], warnings, counts)


def supplier_rows_for(suppliers: list[Supplier]) -> tuple[list[list], list[str]]:
    rows = []
    errors = []
    for supplier in sorted(suppliers, key=lambda item: item.name):
        missing = missing_supplier_fields(supplier)
        if missing:
            errors.append(f"供應商「{supplier.name}」缺少必填欄位：{'、'.join(missing)}。")
            continue
        rows.append([
            supplier.name,
            supplier.owner or "",
            supplier.tax_id or "",
            supplier.address or "",
            supplier.phone or "",
        ])
    return rows, errors


def build_supplier_export(db: Session, branch_id: int, supplier_ids: list[int]) -> ExportResult:
    branch = db.get(Branch, branch_id)
    if not branch:
        return ExportResult({}, ["找不到分店。"], [], {})
    suppliers = visible_items(db, Supplier, "supplier", branch_id, Supplier.name)
    if supplier_ids:
        selected_ids = set(supplier_ids)
        suppliers = [supplier for supplier in suppliers if supplier.id in selected_ids]
    if not suppliers:
        return ExportResult({}, ["沒有選到可用的供應商。"], [], {})
    rows, errors = supplier_rows_for(suppliers)
    if errors:
        return ExportResult({}, errors, [], {})

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = EXPORTS / f"official_excel_{stamp}"
    out_dir.mkdir(parents=True, exist_ok=True)
    supplier_path = copy_template("supplierExcelExample.xlsx", out_dir, f"{safe_filename(branch.name)}_供應商.xlsx")
    count = write_rows(supplier_path, "Data", rows)
    return ExportResult({"供應商": supplier_path}, [], [], {"suppliers": count})


def build_seasoning_export(
    db: Session,
    branch_id: int,
    start: date,
    end: date,
    seasoning_ids: list[int],
) -> ExportResult:
    branch = db.get(Branch, branch_id)
    if not branch:
        return ExportResult({}, ["找不到分店。"], [], {})
    if end < start:
        return ExportResult({}, ["結束日期不能早於開始日期。"], [], {})

    suppliers = {supplier.id: supplier for supplier in visible_items(db, Supplier, "supplier", branch_id, Supplier.name)}
    seasonings = visible_items(db, Seasoning, "seasoning", branch_id, Seasoning.name)
    if seasoning_ids:
        selected_ids = set(seasoning_ids)
        seasonings = [seasoning for seasoning in seasonings if seasoning.id in selected_ids]
    if not seasonings:
        return ExportResult({}, ["沒有選到可用的調味料。"], [], {})

    errors: list[str] = []
    rows: list[list] = []
    for seasoning in seasonings:
        supplier = suppliers.get(seasoning.supplier_id)
        if not supplier:
            errors.append(f"調味料「{seasoning.name}」沒有啟用中的供應商。")
            continue
        purchase_date = purchase_date_for(start, supplier)
        rows.append([
            branch.school_name,
            branch.service_location,
            branch.restaurant_name,
            seasoning.name,
            purchase_date,
            "",
            "",
            start,
            end,
            supplier.name,
            seasoning.manufacturer or "",
            seasoning.certification or "",
            seasoning.certification_no or "",
            "",
            seasoning.product_name or "",
            "",
            seasoning.unit or "",
            seasoning.non_gmo_soy or "",
            seasoning.non_gmo_corn or "",
            seasoning.processed or "",
        ])
    if errors:
        return ExportResult({}, errors, [], {})

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = EXPORTS / f"official_excel_{stamp}"
    out_dir.mkdir(parents=True, exist_ok=True)
    seasoning_path = copy_template("seasoningstockdataCollegeExcelExample.xlsx", out_dir, f"{safe_filename(branch.name)}_調味料_{start:%Y%m%d}_{end:%Y%m%d}.xlsx")
    count = write_rows(seasoning_path, "Sheet1", rows, {5, 6, 7, 8, 9})
    return ExportResult({"調味料": seasoning_path}, [], [], {"seasonings": count})
