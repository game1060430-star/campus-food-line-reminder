import os
import secrets
from io import BytesIO
from datetime import date, timedelta
from pathlib import Path
from fastapi import FastAPI, Depends, Request, Form, HTTPException, UploadFile, File
from fastapi.responses import FileResponse, HTMLResponse, PlainTextResponse, RedirectResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from sqlalchemy import inspect, or_, select, text
from .database import Base, engine, get_db
from .line_bot import LineConfig, handle_line_event, reply_text, verify_line_signature
from .models import Branch, Supplier, Ingredient, Seasoning, Recipe, RecipeIngredient, RecipeSeasoning, DailyMenu, ItemBranchScope, LineBindingCode, UploadConfirmation
from .services.excel_export import EXPORTS, build_export, build_seasoning_export, build_supplier_export, parse_excluded_dates, service_dates
from .services.upload_reminders import send_due_upload_reminders

Base.metadata.create_all(bind=engine)

def ensure_schema():
    inspector = inspect(engine)
    if "line_binding_codes" in inspector.get_table_names():
        columns = {c["name"] for c in inspector.get_columns("line_binding_codes")}
        if "role" not in columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE line_binding_codes ADD COLUMN role VARCHAR(30) DEFAULT 'operator'"))
    if "suppliers" in inspector.get_table_names():
        columns = {c["name"] for c in inspector.get_columns("suppliers")}
        if "delivery_weekdays" not in columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE suppliers ADD COLUMN delivery_weekdays VARCHAR(20)"))
    if "ingredients" in inspector.get_table_names():
        columns = {c["name"]: c for c in inspector.get_columns("ingredients")}
        needs_rebuild = any(columns.get(name, {}).get("nullable") is False for name in ("product_name", "origin", "supplier_id"))
        if needs_rebuild and engine.dialect.name == "sqlite":
            with engine.begin() as conn:
                conn.execute(text("PRAGMA foreign_keys=OFF"))
                conn.execute(text("""
                    CREATE TABLE ingredients_new (
                        id INTEGER NOT NULL PRIMARY KEY,
                        branch_id INTEGER,
                        product_name VARCHAR(200),
                        ingredient_name VARCHAR(200) NOT NULL,
                        origin VARCHAR(120),
                        supplier_id INTEGER,
                        manufacturer VARCHAR(200),
                        processed_food VARCHAR(1),
                        non_gmo_soy VARCHAR(1),
                        non_gmo_corn VARCHAR(1),
                        certification VARCHAR(200),
                        certification_no VARCHAR(120),
                        unit VARCHAR(50),
                        calories VARCHAR(50),
                        active BOOLEAN,
                        FOREIGN KEY(branch_id) REFERENCES branches (id),
                        FOREIGN KEY(supplier_id) REFERENCES suppliers (id)
                    )
                """))
                conn.execute(text("""
                    INSERT INTO ingredients_new (
                        id, branch_id, product_name, ingredient_name, origin, supplier_id,
                        manufacturer, processed_food, non_gmo_soy, non_gmo_corn,
                        certification, certification_no, unit, calories, active
                    )
                    SELECT
                        id, branch_id, product_name, ingredient_name, origin, supplier_id,
                        manufacturer, processed_food, non_gmo_soy, non_gmo_corn,
                        certification, certification_no, unit, calories, active
                    FROM ingredients
                """))
                conn.execute(text("DROP TABLE ingredients"))
                conn.execute(text("ALTER TABLE ingredients_new RENAME TO ingredients"))
                conn.execute(text("PRAGMA foreign_keys=ON"))

ensure_schema()
app = FastAPI(title="校園食材登錄助手")
BASE = Path(__file__).resolve().parent
ROOT_DIR = BASE.parent
templates = Jinja2Templates(directory=str(BASE/"templates"))
app.mount("/static", StaticFiles(directory=str(BASE/"static")), name="static")
app.mount("/phone-app", StaticFiles(directory=str(ROOT_DIR/"phone_app"), html=True), name="phone_app")

PUBLIC_PREFIXES=("/api/line/webhook","/static/","/phone-app","/health","/login")
DEFAULT_WEEKDAYS="0,1,2,3,4"

def web_access_token() -> str:
    return os.getenv("WEB_ACCESS_TOKEN","").strip()

def scoped_query(model, branch_id:int|None):
    if branch_id:
        return or_(model.branch_id==branch_id, model.branch_id.is_(None))
    return True

ITEM_TYPES={Supplier:"supplier", Ingredient:"ingredient", Seasoning:"seasoning", Recipe:"recipe"}

def optional_branch_id(raw:str) -> int|None:
    return int(raw) if raw else None

def redirect_with_branch(path:str, branch_id:int|None):
    suffix=f"?branch_id={branch_id}" if branch_id else ""
    return RedirectResponse(f"{path}{suffix}",303)

def update_scope(item, branch_id:str, path:str, db:Session):
    branch_value=optional_branch_id(branch_id)
    item.branch_id=branch_value
    db.commit()
    return redirect_with_branch(path,branch_value)

def explicit_scope_ids(db:Session, item_type:str, item_id:int) -> set[int]:
    return set(db.scalars(select(ItemBranchScope.branch_id).where(ItemBranchScope.item_type==item_type, ItemBranchScope.item_id==item_id)).all())

def scope_ids_for(item, item_type:str, branches:list[Branch], db:Session) -> set[int]:
    explicit=explicit_scope_ids(db,item_type,item.id)
    if explicit:
        return explicit
    if item.branch_id:
        return {item.branch_id}
    return {b.id for b in branches}

def is_available_for_branch(item, item_type:str, branch_id:int|None, branches:list[Branch], db:Session) -> bool:
    if not branch_id:
        return True
    return branch_id in scope_ids_for(item,item_type,branches,db)

def filter_available(items:list, item_type:str, branch_id:int|None, branches:list[Branch], db:Session) -> list:
    return [item for item in items if is_available_for_branch(item,item_type,branch_id,branches,db)]

def scope_map_for(items:list, item_type:str, branches:list[Branch], db:Session) -> dict[int,set[int]]:
    return {item.id: scope_ids_for(item,item_type,branches,db) for item in items}

def set_item_scopes(item, item_type:str, branch_ids:list[int], db:Session):
    for old in db.scalars(select(ItemBranchScope).where(ItemBranchScope.item_type==item_type, ItemBranchScope.item_id==item.id)).all():
        db.delete(old)
    all_branch_ids={b.id for b in db.scalars(select(Branch).where(Branch.active==True)).all()}
    if set(branch_ids) == all_branch_ids:
        branch_ids=[]
    item.branch_id=None if not branch_ids else branch_ids[0]
    for branch_id in sorted(set(branch_ids)):
        db.add(ItemBranchScope(item_type=item_type,item_id=item.id,branch_id=branch_id))

def parse_weekdays(raw:str) -> set[int]:
    return {int(x) for x in (raw or DEFAULT_WEEKDAYS).split(",") if x}

def join_weekdays(values:list[str]) -> str|None:
    days=sorted({int(value) for value in values if str(value).isdigit()})
    return ",".join(str(day) for day in days) if days else None

def parse_bulk_lines(raw:str) -> list[str]:
    seen=set()
    result=[]
    for line in raw.replace("，","\n").replace(",","\n").splitlines():
        name=line.strip()
        if name and name not in seen:
            seen.add(name)
            result.append(name)
    return result

def clean_text(value) -> str:
    return str(value or "").strip()

def find_supplier_by_name(db:Session, name:str) -> Supplier|None:
    if not name:
        return None
    return db.scalar(select(Supplier).where(Supplier.active==True, Supplier.name==name))

def create_or_update_supplier(db:Session, branch_id:int|None, name:str, owner:str="", tax_id:str="", address:str="", phone:str="") -> bool:
    supplier=find_supplier_by_name(db,name)
    if supplier:
        supplier.owner=owner or supplier.owner
        supplier.tax_id=tax_id or supplier.tax_id
        supplier.address=address or supplier.address
        supplier.phone=phone or supplier.phone
        return False
    db.add(Supplier(name=name,branch_id=branch_id,owner=owner or None,tax_id=tax_id or None,address=address or None,phone=phone or None))
    return True

def read_uploaded_workbook(file:UploadFile):
    return load_workbook(BytesIO(file.file.read()), data_only=True)

def import_supplier_workbook(db:Session, branch_id:int|None, file:UploadFile) -> tuple[int,list[str]]:
    wb=read_uploaded_workbook(file)
    ws=wb["Data"] if "Data" in wb.sheetnames else wb[wb.sheetnames[0]]
    created=0
    errors=[]
    for row_index,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        name=clean_text(row[0] if len(row)>0 else "")
        if not name:
            continue
        owner=clean_text(row[1] if len(row)>1 else "")
        tax_id=clean_text(row[2] if len(row)>2 else "")
        address=clean_text(row[3] if len(row)>3 else "")
        phone=clean_text(row[4] if len(row)>4 else "")
        if not all([owner,tax_id,address,phone]):
            errors.append(f"第 {row_index} 列供應商「{name}」缺少負責人、統編、地址或電話，已略過。")
            continue
        if create_or_update_supplier(db,branch_id,name,owner,tax_id,address,phone):
            created+=1
    db.commit()
    return created,errors

def import_ingredient_workbook(db:Session, branch_id:int|None, file:UploadFile) -> tuple[int,list[str]]:
    wb=read_uploaded_workbook(file)
    ws=wb[wb.sheetnames[0]]
    created=0
    errors=[]
    existing={item.ingredient_name for item in db.scalars(select(Ingredient).where(Ingredient.active==True)).all() if item.branch_id in (None,branch_id)}
    for row_index,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        product_name=clean_text(row[5] if len(row)>5 else "")
        ingredient_name=clean_text(row[6] if len(row)>6 else "") or product_name
        origin=clean_text(row[7] if len(row)>7 else "")
        supplier_name=clean_text(row[8] if len(row)>8 else "")
        if not ingredient_name:
            continue
        if ingredient_name in existing:
            continue
        supplier=find_supplier_by_name(db,supplier_name)
        if supplier_name and not supplier:
            errors.append(f"第 {row_index} 列食材「{ingredient_name}」找不到供應商「{supplier_name}」，先建立成待補供應商。")
        db.add(Ingredient(product_name=product_name or ingredient_name,ingredient_name=ingredient_name,origin=origin or "臺灣",supplier_id=supplier.id if supplier else None,branch_id=branch_id))
        existing.add(ingredient_name)
        created+=1
    db.commit()
    return created,errors

def import_recipe_workbook(db:Session, branch_id:int|None, file:UploadFile) -> tuple[int,list[str]]:
    wb=read_uploaded_workbook(file)
    ws=wb[wb.sheetnames[0]]
    created=0
    existing={item.name for item in db.scalars(select(Recipe).where(Recipe.active==True)).all() if item.branch_id in (None,branch_id)}
    for row in ws.iter_rows(min_row=2,values_only=True):
        name=clean_text(row[5] if len(row)>5 else "")
        if not name or name in existing:
            continue
        calories_raw=row[7] if len(row)>7 else 0
        try:
            calories=int(calories_raw or 0)
        except (TypeError,ValueError):
            calories=0
        db.add(Recipe(name=name,category=None,calories=calories,branch_id=branch_id))
        existing.add(name)
        created+=1
    db.commit()
    return created,[]

def import_seasoning_workbook(db:Session, branch_id:int|None, file:UploadFile) -> tuple[int,list[str]]:
    wb=read_uploaded_workbook(file)
    ws=wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
    created=0
    errors=[]
    existing={item.name for item in db.scalars(select(Seasoning).where(Seasoning.active==True)).all() if item.branch_id in (None,branch_id)}
    for row_index,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        name=clean_text(row[3] if len(row)>3 else "")
        supplier_name=clean_text(row[9] if len(row)>9 else "")
        if not name or name in existing:
            continue
        supplier=find_supplier_by_name(db,supplier_name)
        if not supplier:
            errors.append(f"第 {row_index} 列調味料「{name}」找不到供應商「{supplier_name}」，已略過。")
            continue
        db.add(Seasoning(name=name,supplier_id=supplier.id,branch_id=branch_id))
        existing.add(name)
        created+=1
    db.commit()
    return created,errors

def confirm_upload_dates(db:Session, branch_id:int, dates:list[date], confirmed_by:str="", note:str="") -> int:
    count=0
    for service_date in dates:
        existing=db.scalar(select(UploadConfirmation).where(UploadConfirmation.branch_id==branch_id, UploadConfirmation.service_date==service_date))
        if existing:
            existing.confirmed_by=confirmed_by or existing.confirmed_by
            existing.note=note or existing.note
        else:
            db.add(UploadConfirmation(branch_id=branch_id,service_date=service_date,confirmed_by=confirmed_by or None,note=note or None))
            count+=1
    db.commit()
    return count

@app.middleware("http")
async def require_web_access(request:Request, call_next):
    path=request.url.path
    if path.startswith(PUBLIC_PREFIXES):
        return await call_next(request)
    token=web_access_token()
    query_token=request.query_params.get("token","")
    if token and query_token and secrets.compare_digest(query_token, token):
        return await call_next(request)
    if token and request.cookies.get("web_access") != token:
        return PlainTextResponse("此頁面只開放給已綁定的 LINE 管理員。請從 LINE 官方帳號取得安全連結。",status_code=403)
    return await call_next(request)

@app.get("/login")
def login_with_token(token:str="", next:str="/"):
    if not web_access_token() or not secrets.compare_digest(token, web_access_token()):
        raise HTTPException(403)
    if not next.startswith("/") or next.startswith("//"):
        next="/"
    response=RedirectResponse(next,303)
    response.set_cookie("web_access",token,httponly=True,samesite="lax",max_age=60*60*24*30)
    return response

@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, db: Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    counts={
        "suppliers": len(db.scalars(select(Supplier).where(Supplier.active==True)).all()),
        "ingredients": len(db.scalars(select(Ingredient).where(Ingredient.active==True)).all()),
        "seasonings": len(db.scalars(select(Seasoning).where(Seasoning.active==True)).all()),
        "recipes": len(db.scalars(select(Recipe).where(Recipe.active==True)).all()),
    }
    return templates.TemplateResponse("dashboard.html",{"request":request,"branches":branches,"counts":counts})

@app.get("/dashboard", response_class=HTMLResponse)
def dashboard_alias(request: Request, db: Session=Depends(get_db)):
    return dashboard(request,db)

@app.get("/branches", response_class=HTMLResponse)
def branches_page(request:Request, db:Session=Depends(get_db)):
    items=db.scalars(select(Branch)).all()
    codes={c.branch_id:c for c in db.scalars(select(LineBindingCode).where(LineBindingCode.active==True).order_by(LineBindingCode.id.desc())).all()}
    return templates.TemplateResponse("branches.html",{"request":request,"items":items,"codes":codes})

@app.get("/data", response_class=HTMLResponse)
def data_page(request:Request, db:Session=Depends(get_db)):
    counts={
        "suppliers": len(db.scalars(select(Supplier).where(Supplier.active==True)).all()),
        "ingredients": len(db.scalars(select(Ingredient).where(Ingredient.active==True)).all()),
        "seasonings": len(db.scalars(select(Seasoning).where(Seasoning.active==True)).all()),
        "recipes": len(db.scalars(select(Recipe).where(Recipe.active==True)).all()),
    }
    return templates.TemplateResponse("data.html",{"request":request,"counts":counts})

@app.get("/imports", response_class=HTMLResponse)
def imports_page(request:Request, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    selected_branch_id=int(request.query_params.get("branch_id") or 0) or None
    return templates.TemplateResponse("imports.html",{"request":request,"branches":branches,"selected_branch_id":selected_branch_id})

@app.post("/imports", response_class=HTMLResponse)
async def import_existing_file(request:Request, branch_id:str=Form(""), import_type:str=Form(...), file:UploadFile=File(...), db:Session=Depends(get_db)):
    branch_value=optional_branch_id(branch_id)
    errors=[]
    created=0
    if import_type=="suppliers":
        created,errors=import_supplier_workbook(db,branch_value,file)
    elif import_type=="ingredients":
        created,errors=import_ingredient_workbook(db,branch_value,file)
    elif import_type=="recipes":
        created,errors=import_recipe_workbook(db,branch_value,file)
    elif import_type=="seasonings":
        created,errors=import_seasoning_workbook(db,branch_value,file)
    else:
        raise HTTPException(400)
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    return templates.TemplateResponse("imports.html",{"request":request,"branches":branches,"selected_branch_id":branch_value,"created":created,"errors":errors})

@app.get("/bulk-setup", response_class=HTMLResponse)
def bulk_setup_page(request:Request, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    selected_branch_id=int(request.query_params.get("branch_id") or 0) or None
    return templates.TemplateResponse("bulk_setup.html",{"request":request,"branches":branches,"selected_branch_id":selected_branch_id})

@app.post("/bulk-setup/ingredients", response_class=HTMLResponse)
async def bulk_create_ingredients(request:Request, db:Session=Depends(get_db)):
    form=await request.form()
    branch_id=optional_branch_id(str(form.get("branch_id") or ""))
    names=parse_bulk_lines(str(form.get("names") or ""))
    existing={
        item.ingredient_name
        for item in db.scalars(select(Ingredient).where(Ingredient.active==True)).all()
        if item.branch_id in (None, branch_id)
    }
    created=0
    skipped=[]
    for name in names:
        if name in existing:
            skipped.append(name)
            continue
        db.add(Ingredient(product_name=name,ingredient_name=name,origin="臺灣",supplier_id=None,branch_id=branch_id))
        existing.add(name)
        created+=1
    db.commit()
    return RedirectResponse(f"/bulk-setup?branch_id={branch_id or ''}&notice=ingredients&created={created}&skipped={len(skipped)}",303)

@app.post("/bulk-setup/recipes", response_class=HTMLResponse)
async def bulk_create_recipes(request:Request, db:Session=Depends(get_db)):
    form=await request.form()
    branch_id=optional_branch_id(str(form.get("branch_id") or ""))
    names=parse_bulk_lines(str(form.get("names") or ""))
    existing={
        item.name
        for item in db.scalars(select(Recipe).where(Recipe.active==True)).all()
        if item.branch_id in (None, branch_id)
    }
    created=0
    skipped=[]
    for name in names:
        if name in existing:
            skipped.append(name)
            continue
        db.add(Recipe(name=name,category=None,calories=0,branch_id=branch_id))
        existing.add(name)
        created+=1
    db.commit()
    return RedirectResponse(f"/bulk-setup?branch_id={branch_id or ''}&notice=recipes&created={created}&skipped={len(skipped)}",303)

@app.get("/exports", response_class=HTMLResponse)
def exports_page(request:Request, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    selected_branch_id=int(request.query_params.get("branch_id") or (branches[0].id if branches else 0))
    recipes_all=db.scalars(select(Recipe).where(Recipe.active==True).order_by(Recipe.name)).all()
    recipes=filter_available(recipes_all,"recipe",selected_branch_id,branches,db)
    return templates.TemplateResponse("exports.html",{"request":request,"branches":branches,"recipes":recipes,"selected_branch_id":selected_branch_id})

@app.post("/exports")
async def create_export(request:Request, db:Session=Depends(get_db)):
    form=await request.form()
    branch_id=int(form.get("branch_id") or 0)
    start_date=date.fromisoformat(str(form.get("start_date")))
    end_date=date.fromisoformat(str(form.get("end_date")))
    weekdays={int(x) for x in str(form.get("weekdays") or "0,1,2,3,4").split(",") if x}
    excluded=parse_excluded_dates(str(form.get("excluded_dates") or ""))
    recipe_ids=[int(x) for x in form.getlist("recipe_ids") if str(x).isdigit()]
    file_types={str(x) for x in form.getlist("file_types")}
    if not file_types:
        file_types={"menu","ingredients"}
    result=build_export(db,branch_id,start_date,end_date,weekdays,excluded,recipe_ids,file_types)
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    recipes_all=db.scalars(select(Recipe).where(Recipe.active==True).order_by(Recipe.name)).all()
    recipes=filter_available(recipes_all,"recipe",branch_id,branches,db)
    if result.errors:
        return templates.TemplateResponse("exports.html",{"request":request,"branches":branches,"recipes":recipes,"errors":result.errors,"warnings":result.warnings,"selected_branch_id":branch_id})
    download_files={label: path.relative_to(EXPORTS).as_posix() for label,path in result.files.items()}
    confirm_context={
        "branch_id": branch_id,
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "weekdays": ",".join(str(x) for x in sorted(weekdays)),
        "excluded_dates": ",".join(sorted(d.isoformat() for d in excluded)),
    }
    return templates.TemplateResponse("exports.html",{"request":request,"branches":branches,"recipes":recipes,"warnings":result.warnings,"download_files":download_files,"download_token":web_access_token(),"row_counts":result.row_counts,"selected_branch_id":branch_id,"confirm_context":confirm_context})

@app.get("/master-exports", response_class=HTMLResponse)
def master_exports_page(request:Request, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    selected_branch_id=int(request.query_params.get("branch_id") or (branches[0].id if branches else 0))
    suppliers=filter_available(db.scalars(select(Supplier).where(Supplier.active==True).order_by(Supplier.name)).all(),"supplier",selected_branch_id,branches,db)
    seasonings=filter_available(db.scalars(select(Seasoning).where(Seasoning.active==True).order_by(Seasoning.name)).all(),"seasoning",selected_branch_id,branches,db)
    return templates.TemplateResponse("master_exports.html",{"request":request,"branches":branches,"selected_branch_id":selected_branch_id,"suppliers":suppliers,"seasonings":seasonings})

@app.post("/master-exports/suppliers", response_class=HTMLResponse)
async def create_supplier_master_export(request:Request, db:Session=Depends(get_db)):
    form=await request.form()
    branch_id=int(form.get("branch_id") or 0)
    supplier_ids=[int(x) for x in form.getlist("supplier_ids") if str(x).isdigit()]
    result=build_supplier_export(db,branch_id,supplier_ids)
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    suppliers=filter_available(db.scalars(select(Supplier).where(Supplier.active==True).order_by(Supplier.name)).all(),"supplier",branch_id,branches,db)
    seasonings=filter_available(db.scalars(select(Seasoning).where(Seasoning.active==True).order_by(Seasoning.name)).all(),"seasoning",branch_id,branches,db)
    download_files={label: path.relative_to(EXPORTS).as_posix() for label,path in result.files.items()}
    return templates.TemplateResponse("master_exports.html",{"request":request,"branches":branches,"selected_branch_id":branch_id,"suppliers":suppliers,"seasonings":seasonings,"errors":result.errors,"warnings":result.warnings,"download_files":download_files,"download_token":web_access_token(),"row_counts":result.row_counts})

@app.post("/master-exports/seasonings", response_class=HTMLResponse)
async def create_seasoning_master_export(request:Request, db:Session=Depends(get_db)):
    form=await request.form()
    branch_id=int(form.get("branch_id") or 0)
    start_date=date.fromisoformat(str(form.get("start_date")))
    end_date=date.fromisoformat(str(form.get("end_date")))
    seasoning_ids=[int(x) for x in form.getlist("seasoning_ids") if str(x).isdigit()]
    result=build_seasoning_export(db,branch_id,start_date,end_date,seasoning_ids)
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    suppliers=filter_available(db.scalars(select(Supplier).where(Supplier.active==True).order_by(Supplier.name)).all(),"supplier",branch_id,branches,db)
    seasonings=filter_available(db.scalars(select(Seasoning).where(Seasoning.active==True).order_by(Seasoning.name)).all(),"seasoning",branch_id,branches,db)
    download_files={label: path.relative_to(EXPORTS).as_posix() for label,path in result.files.items()}
    return templates.TemplateResponse("master_exports.html",{"request":request,"branches":branches,"selected_branch_id":branch_id,"suppliers":suppliers,"seasonings":seasonings,"errors":result.errors,"warnings":result.warnings,"download_files":download_files,"download_token":web_access_token(),"row_counts":result.row_counts})

@app.get("/uploads", response_class=HTMLResponse)
def uploads_page(request:Request, branch_id:int|None=None, date:str|None=None, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    selected_branch_id=branch_id or (branches[0].id if branches else 0)
    selected_date=date or date_today_iso()
    start=date_from_iso(selected_date)
    days=[start+timedelta(days=i) for i in range(14)]
    confirmations={
        (c.branch_id,c.service_date): c
        for c in db.scalars(select(UploadConfirmation).where(UploadConfirmation.service_date.in_(days))).all()
    }
    menu_counts={}
    for item in db.scalars(select(DailyMenu).where(DailyMenu.service_date.in_(days))).all():
        menu_counts[(item.branch_id,item.service_date)]=menu_counts.get((item.branch_id,item.service_date),0)+1
    return templates.TemplateResponse("uploads.html",{"request":request,"branches":branches,"selected_branch_id":selected_branch_id,"selected_date":selected_date,"days":days,"confirmations":confirmations,"menu_counts":menu_counts,"notice":request.query_params.get("notice")})

def date_today_iso() -> str:
    return date.today().isoformat()

def date_from_iso(value:str) -> date:
    return date.fromisoformat(value)

@app.post("/uploads/confirm")
def confirm_upload(branch_id:int=Form(...), start_date:date=Form(...), end_date:date=Form(...), weekdays:str=Form(DEFAULT_WEEKDAYS), excluded_dates:str=Form(""), confirmed_by:str=Form(""), note:str=Form(""), db:Session=Depends(get_db)):
    excluded=parse_excluded_dates(excluded_dates)
    dates=service_dates(start_date,end_date,parse_weekdays(weekdays),excluded)
    confirm_upload_dates(db,branch_id,dates,confirmed_by,note)
    return RedirectResponse(f"/uploads?branch_id={branch_id}&date={start_date.isoformat()}&notice=confirmed",303)

@app.get("/exports/files/{file_path:path}")
def download_export_file(file_path:str):
    target=(EXPORTS/file_path).resolve()
    exports_root=EXPORTS.resolve()
    if exports_root not in target.parents or not target.exists():
        raise HTTPException(404)
    return FileResponse(target,filename=target.name)

@app.post("/branches")
def create_branch(name:str=Form(...), school_name:str=Form(...), service_location:str=Form(...), restaurant_name:str=Form(...), db:Session=Depends(get_db)):
    b=Branch(name=name,school_name=school_name,service_location=service_location,restaurant_name=restaurant_name)
    db.add(b); db.commit()
    return RedirectResponse('/branches',303)

@app.post("/branches/{branch_id}/line-code")
def create_line_binding_code(branch_id:int, db:Session=Depends(get_db)):
    branch=db.get(Branch,branch_id)
    if not branch: raise HTTPException(404)
    for old in db.scalars(select(LineBindingCode).where(LineBindingCode.branch_id==branch_id,LineBindingCode.active==True)).all():
        old.active=False
    code=secrets.token_hex(3).upper()
    while db.scalar(select(LineBindingCode).where(LineBindingCode.code==code)):
        code=secrets.token_hex(3).upper()
    db.add(LineBindingCode(branch_id=branch_id,code=code,role="operator"))
    db.commit()
    return RedirectResponse('/branches',303)

@app.get("/recipes", response_class=HTMLResponse)
def recipes_page(request:Request, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    branch_map={b.id:b for b in branches}
    selected_branch_id=int(request.query_params.get("branch_id") or 0) or None
    recipes=filter_available(db.scalars(select(Recipe).where(Recipe.active==True).order_by(Recipe.name)).all(),"recipe",selected_branch_id,branches,db)
    ingredients=filter_available(db.scalars(select(Ingredient).where(Ingredient.active==True).order_by(Ingredient.ingredient_name)).all(),"ingredient",selected_branch_id,branches,db)
    seasonings=filter_available(db.scalars(select(Seasoning).where(Seasoning.active==True).order_by(Seasoning.name)).all(),"seasoning",selected_branch_id,branches,db)
    recipe_scope_map=scope_map_for(recipes,"recipe",branches,db)
    recipe_ingredients=db.scalars(select(RecipeIngredient)).all()
    recipe_seasonings=db.scalars(select(RecipeSeasoning)).all()
    ingredient_map={i.id:i for i in ingredients}
    seasoning_map={s.id:s for s in seasonings}
    ingredients_by_recipe={}
    seasonings_by_recipe={}
    ingredient_ids_by_recipe={}
    seasoning_ids_by_recipe={}
    for item in recipe_ingredients:
        if item.ingredient_id in ingredient_map:
            ingredients_by_recipe.setdefault(item.recipe_id,[]).append(ingredient_map[item.ingredient_id].ingredient_name)
            ingredient_ids_by_recipe.setdefault(item.recipe_id,[]).append(item.ingredient_id)
    for item in recipe_seasonings:
        if item.seasoning_id in seasoning_map:
            seasonings_by_recipe.setdefault(item.recipe_id,[]).append(seasoning_map[item.seasoning_id].name)
            seasoning_ids_by_recipe.setdefault(item.recipe_id,[]).append(item.seasoning_id)
    return templates.TemplateResponse("recipes.html",{"request":request,"recipes":recipes,"branches":branches,"branch_map":branch_map,"ingredients":ingredients,"seasonings":seasonings,"ingredients_by_recipe":ingredients_by_recipe,"seasonings_by_recipe":seasonings_by_recipe,"ingredient_ids_by_recipe":ingredient_ids_by_recipe,"seasoning_ids_by_recipe":seasoning_ids_by_recipe,"scope_map":recipe_scope_map,"selected_branch_id":selected_branch_id})

@app.post("/recipes")
async def create_recipe(request:Request, db:Session=Depends(get_db)):
    form=await request.form()
    name=str(form.get("name") or "").strip()
    if not name: raise HTTPException(400)
    category=str(form.get("category") or "").strip()
    calories=int(form.get("calories") or 0)
    branch_id=optional_branch_id(str(form.get("branch_id") or ""))
    r=Recipe(name=name,category=category or None,calories=calories,branch_id=branch_id)
    db.add(r); db.flush()
    for ingredient_id in form.getlist("ingredient_ids"):
        if str(ingredient_id).isdigit():
            db.add(RecipeIngredient(recipe_id=r.id,ingredient_id=int(ingredient_id)))
    for seasoning_id in form.getlist("seasoning_ids"):
        if str(seasoning_id).isdigit():
            db.add(RecipeSeasoning(recipe_id=r.id,seasoning_id=int(seasoning_id)))
    db.commit()
    return redirect_with_branch('/recipes',branch_id)

@app.post("/recipes/{recipe_id}/update")
async def update_recipe(recipe_id:int, request:Request, db:Session=Depends(get_db)):
    recipe=db.get(Recipe,recipe_id)
    if not recipe: raise HTTPException(404)
    form=await request.form()
    name=str(form.get("name") or "").strip()
    if not name: raise HTTPException(400)
    branch_id=optional_branch_id(str(form.get("branch_id") or ""))
    recipe.name=name
    recipe.category=str(form.get("category") or "").strip() or None
    recipe.calories=int(form.get("calories") or 0)
    recipe.branch_id=branch_id
    for link in db.scalars(select(RecipeIngredient).where(RecipeIngredient.recipe_id==recipe.id)).all():
        db.delete(link)
    for link in db.scalars(select(RecipeSeasoning).where(RecipeSeasoning.recipe_id==recipe.id)).all():
        db.delete(link)
    db.flush()
    for ingredient_id in form.getlist("ingredient_ids"):
        if str(ingredient_id).isdigit():
            db.add(RecipeIngredient(recipe_id=recipe.id,ingredient_id=int(ingredient_id)))
    for seasoning_id in form.getlist("seasoning_ids"):
        if str(seasoning_id).isdigit():
            db.add(RecipeSeasoning(recipe_id=recipe.id,seasoning_id=int(seasoning_id)))
    db.commit()
    return redirect_with_branch('/recipes',branch_id)

@app.post("/recipes/{recipe_id}/delete")
def delete_recipe(recipe_id:int, db:Session=Depends(get_db)):
    recipe=db.get(Recipe,recipe_id)
    if not recipe: raise HTTPException(404)
    recipe.active=False
    db.commit()
    return RedirectResponse('/recipes',303)

@app.post("/recipes/{recipe_id}/scope")
async def update_recipe_scope(recipe_id:int, request:Request, db:Session=Depends(get_db)):
    recipe=db.get(Recipe,recipe_id)
    if not recipe: raise HTTPException(404)
    form=await request.form()
    branch_ids=[int(x) for x in form.getlist("branch_ids") if str(x).isdigit()]
    set_item_scopes(recipe,"recipe",branch_ids,db)
    db.commit()
    return redirect_with_branch('/recipes',branch_ids[0] if branch_ids else None)

@app.get("/suppliers", response_class=HTMLResponse)
def suppliers_page(request:Request, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    selected_branch_id=int(request.query_params.get("branch_id") or 0) or None
    items=filter_available(db.scalars(select(Supplier).where(Supplier.active==True).order_by(Supplier.name)).all(),"supplier",selected_branch_id,branches,db)
    branch_map={b.id:b for b in branches}
    supplier_scope_map=scope_map_for(items,"supplier",branches,db)
    assignable_ingredients=filter_available(db.scalars(select(Ingredient).where(Ingredient.active==True).order_by(Ingredient.ingredient_name)).all(),"ingredient",selected_branch_id,branches,db)
    ingredient_assign_options={supplier.id:[ingredient for ingredient in assignable_ingredients if ingredient.supplier_id != supplier.id] for supplier in items}
    supplier_dependencies={}
    for supplier in items:
        ingredients=db.scalars(select(Ingredient).where(Ingredient.supplier_id==supplier.id,Ingredient.active==True).order_by(Ingredient.ingredient_name)).all()
        seasonings=db.scalars(select(Seasoning).where(Seasoning.supplier_id==supplier.id,Seasoning.active==True).order_by(Seasoning.name)).all()
        supplier_dependencies[supplier.id]={"ingredients":ingredients,"seasonings":seasonings}
    weekdays=[(0,"週一"),(1,"週二"),(2,"週三"),(3,"週四"),(4,"週五"),(5,"週六"),(6,"週日")]
    return templates.TemplateResponse("suppliers.html",{"request":request,"branches":branches,"items":items,"branch_map":branch_map,"supplier_dependencies":supplier_dependencies,"ingredient_assign_options":ingredient_assign_options,"scope_map":supplier_scope_map,"weekdays":weekdays,"blocked":request.query_params.get("blocked"),"selected_branch_id":selected_branch_id})

@app.post("/suppliers")
async def create_supplier(request:Request, name:str=Form(...), branch_id:str=Form(""), owner:str=Form(...), tax_id:str=Form(...), address:str=Form(...), phone:str=Form(...), db:Session=Depends(get_db)):
    if not all(value.strip() for value in [name, owner, tax_id, address, phone]):
        raise HTTPException(400, "供應商名稱、負責人、統編、地址、電話都是必填。")
    form=await request.form()
    branch_value=optional_branch_id(branch_id)
    db.add(Supplier(name=name.strip(),branch_id=branch_value,owner=owner.strip(),tax_id=tax_id.strip(),address=address.strip(),phone=phone.strip(),delivery_weekdays=join_weekdays(form.getlist("delivery_weekdays"))))
    db.commit()
    return redirect_with_branch('/suppliers',branch_value)

@app.post("/suppliers/{supplier_id}/update")
async def update_supplier(supplier_id:int, request:Request, name:str=Form(...), branch_id:str=Form(""), owner:str=Form(...), tax_id:str=Form(...), address:str=Form(...), phone:str=Form(...), db:Session=Depends(get_db)):
    supplier=db.get(Supplier,supplier_id)
    if not supplier: raise HTTPException(404)
    if not all(value.strip() for value in [name, owner, tax_id, address, phone]):
        raise HTTPException(400, "供應商名稱、負責人、統編、地址、電話都是必填。")
    form=await request.form()
    branch_value=optional_branch_id(branch_id)
    supplier.name=name.strip()
    supplier.branch_id=branch_value
    supplier.owner=owner.strip()
    supplier.tax_id=tax_id.strip()
    supplier.address=address.strip()
    supplier.phone=phone.strip()
    supplier.delivery_weekdays=join_weekdays(form.getlist("delivery_weekdays"))
    db.commit()
    return redirect_with_branch('/suppliers',branch_value)

@app.post("/suppliers/{supplier_id}/delete")
def delete_supplier(supplier_id:int, db:Session=Depends(get_db)):
    supplier=db.get(Supplier,supplier_id)
    if not supplier: raise HTTPException(404)
    active_ingredients=db.scalars(select(Ingredient).where(Ingredient.supplier_id==supplier_id,Ingredient.active==True)).all()
    active_seasonings=db.scalars(select(Seasoning).where(Seasoning.supplier_id==supplier_id,Seasoning.active==True)).all()
    if active_ingredients or active_seasonings:
        return RedirectResponse('/suppliers?blocked=supplier_has_items',303)
    supplier.active=False
    db.commit()
    return RedirectResponse('/suppliers',303)

@app.post("/suppliers/{supplier_id}/scope")
async def update_supplier_scope(supplier_id:int, request:Request, db:Session=Depends(get_db)):
    supplier=db.get(Supplier,supplier_id)
    if not supplier: raise HTTPException(404)
    form=await request.form()
    branch_ids=[int(x) for x in form.getlist("branch_ids") if str(x).isdigit()]
    set_item_scopes(supplier,"supplier",branch_ids,db)
    db.commit()
    return redirect_with_branch('/suppliers',branch_ids[0] if branch_ids else None)

@app.post("/suppliers/{supplier_id}/replace")
def replace_supplier_then_delete(supplier_id:int, replacement_supplier_id:int=Form(...), db:Session=Depends(get_db)):
    supplier=db.get(Supplier,supplier_id)
    replacement=db.get(Supplier,replacement_supplier_id)
    if not supplier or not replacement or supplier.id==replacement.id: raise HTTPException(404)
    for ingredient in db.scalars(select(Ingredient).where(Ingredient.supplier_id==supplier_id,Ingredient.active==True)).all():
        ingredient.supplier_id=replacement.id
    for seasoning in db.scalars(select(Seasoning).where(Seasoning.supplier_id==supplier_id,Seasoning.active==True)).all():
        seasoning.supplier_id=replacement.id
    supplier.active=False
    db.commit()
    return RedirectResponse('/suppliers',303)

@app.post("/suppliers/{supplier_id}/assign-ingredients")
async def assign_ingredients_to_supplier(supplier_id:int, request:Request, db:Session=Depends(get_db)):
    supplier=db.get(Supplier,supplier_id)
    if not supplier: raise HTTPException(404)
    form=await request.form()
    ingredient_ids=[int(x) for x in form.getlist("ingredient_ids") if str(x).isdigit()]
    branch_id=optional_branch_id(str(form.get("branch_id") or ""))
    for ingredient in db.scalars(select(Ingredient).where(Ingredient.id.in_(ingredient_ids),Ingredient.active==True)).all():
        ingredient.supplier_id=supplier.id
    db.commit()
    return redirect_with_branch('/suppliers',branch_id)

@app.get("/ingredients", response_class=HTMLResponse)
def ingredients_page(request:Request, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    selected_branch_id=int(request.query_params.get("branch_id") or 0) or None
    suppliers=filter_available(db.scalars(select(Supplier).where(Supplier.active==True).order_by(Supplier.name)).all(),"supplier",selected_branch_id,branches,db)
    items=filter_available(db.scalars(select(Ingredient).where(Ingredient.active==True).order_by(Ingredient.ingredient_name)).all(),"ingredient",selected_branch_id,branches,db)
    supplier_map={s.id:s for s in suppliers}
    branch_map={b.id:b for b in branches}
    active_recipes={r.id:r for r in db.scalars(select(Recipe).where(Recipe.active==True)).all()}
    ingredient_recipes={}
    for link in db.scalars(select(RecipeIngredient)).all():
        recipe=active_recipes.get(link.recipe_id)
        if recipe:
            ingredient_recipes.setdefault(link.ingredient_id,[]).append(recipe)
    ingredient_scope_map=scope_map_for(items,"ingredient",branches,db)
    missing_fields_by_ingredient={
        item.id: [
            label for label,value in [
                ("產品名稱", item.product_name),
                ("原產地", item.origin),
                ("供應商", supplier_map.get(item.supplier_id)),
            ] if not value
        ]
        for item in items
    }
    return templates.TemplateResponse("ingredients.html",{"request":request,"branches":branches,"suppliers":suppliers,"items":items,"supplier_map":supplier_map,"branch_map":branch_map,"ingredient_recipes":ingredient_recipes,"scope_map":ingredient_scope_map,"missing_fields_by_ingredient":missing_fields_by_ingredient,"blocked":request.query_params.get("blocked"),"selected_branch_id":selected_branch_id})

@app.post("/ingredients")
def create_ingredient(product_name:str=Form(...), ingredient_name:str=Form(...), origin:str=Form(""), supplier_id:int=Form(...), branch_id:str=Form(""), manufacturer:str=Form(""), unit:str=Form(""), certification:str=Form(""), certification_no:str=Form(""), db:Session=Depends(get_db)):
    branch_value=optional_branch_id(branch_id)
    db.add(Ingredient(product_name=product_name,ingredient_name=ingredient_name,origin=origin.strip() or "臺灣",supplier_id=supplier_id,branch_id=branch_value,manufacturer=manufacturer or None,unit=unit or None,certification=certification or None,certification_no=certification_no or None))
    db.commit()
    return redirect_with_branch('/ingredients',branch_value)

@app.post("/ingredients/bulk-update")
async def bulk_update_ingredients(request:Request, db:Session=Depends(get_db)):
    form=await request.form()
    ids=[int(x) for x in form.getlist("ingredient_ids") if str(x).isdigit()]
    product_names=[str(x).strip() for x in form.getlist("product_names")]
    ingredient_names=[str(x).strip() for x in form.getlist("ingredient_names")]
    origins=[str(x).strip() for x in form.getlist("origins")]
    supplier_ids=[str(x).strip() for x in form.getlist("supplier_ids")]
    branch_ids=[str(x).strip() for x in form.getlist("branch_ids")]
    selected_branch=optional_branch_id(str(form.get("selected_branch_id") or ""))
    for index, ingredient_id in enumerate(ids):
        ingredient=db.get(Ingredient,ingredient_id)
        if not ingredient:
            continue
        name=ingredient_names[index] if index < len(ingredient_names) else ingredient.ingredient_name
        product=product_names[index] if index < len(product_names) else ingredient.product_name
        origin=origins[index] if index < len(origins) else ingredient.origin
        supplier_raw=supplier_ids[index] if index < len(supplier_ids) else ""
        branch_raw=branch_ids[index] if index < len(branch_ids) else ""
        ingredient.ingredient_name=name or ingredient.ingredient_name
        ingredient.product_name=product or ingredient.ingredient_name
        ingredient.origin=origin or "臺灣"
        ingredient.supplier_id=int(supplier_raw) if supplier_raw.isdigit() else None
        ingredient.branch_id=optional_branch_id(branch_raw)
    db.commit()
    return redirect_with_branch('/ingredients',selected_branch)

@app.post("/ingredients/{ingredient_id}/update")
def update_ingredient(ingredient_id:int, product_name:str=Form(...), ingredient_name:str=Form(...), origin:str=Form(""), supplier_id:int=Form(...), branch_id:str=Form(""), manufacturer:str=Form(""), unit:str=Form(""), certification:str=Form(""), certification_no:str=Form(""), db:Session=Depends(get_db)):
    ingredient=db.get(Ingredient,ingredient_id)
    if not ingredient: raise HTTPException(404)
    branch_value=optional_branch_id(branch_id)
    ingredient.product_name=product_name.strip()
    ingredient.ingredient_name=ingredient_name.strip()
    ingredient.origin=origin.strip() or "臺灣"
    ingredient.supplier_id=supplier_id
    ingredient.branch_id=branch_value
    ingredient.manufacturer=manufacturer.strip() or None
    ingredient.unit=unit.strip() or None
    ingredient.certification=certification.strip() or None
    ingredient.certification_no=certification_no.strip() or None
    db.commit()
    return redirect_with_branch('/ingredients',branch_value)

@app.post("/ingredients/{ingredient_id}/delete")
def delete_ingredient(ingredient_id:int, db:Session=Depends(get_db)):
    ingredient=db.get(Ingredient,ingredient_id)
    if not ingredient: raise HTTPException(404)
    active_recipe_ids={r.id for r in db.scalars(select(Recipe).where(Recipe.active==True)).all()}
    used=db.scalars(select(RecipeIngredient).where(RecipeIngredient.ingredient_id==ingredient_id)).all()
    if any(link.recipe_id in active_recipe_ids for link in used):
        return RedirectResponse('/ingredients?blocked=ingredient_used_by_recipe',303)
    ingredient.active=False
    db.commit()
    return RedirectResponse('/ingredients',303)

@app.post("/ingredients/{ingredient_id}/scope")
async def update_ingredient_scope(ingredient_id:int, request:Request, db:Session=Depends(get_db)):
    ingredient=db.get(Ingredient,ingredient_id)
    if not ingredient: raise HTTPException(404)
    form=await request.form()
    branch_ids=[int(x) for x in form.getlist("branch_ids") if str(x).isdigit()]
    set_item_scopes(ingredient,"ingredient",branch_ids,db)
    db.commit()
    return redirect_with_branch('/ingredients',branch_ids[0] if branch_ids else None)

@app.post("/ingredients/{ingredient_id}/replace")
def replace_ingredient_then_delete(ingredient_id:int, replacement_ingredient_id:int=Form(...), db:Session=Depends(get_db)):
    ingredient=db.get(Ingredient,ingredient_id)
    replacement=db.get(Ingredient,replacement_ingredient_id)
    if not ingredient or not replacement or ingredient.id==replacement.id: raise HTTPException(404)
    for link in db.scalars(select(RecipeIngredient).where(RecipeIngredient.ingredient_id==ingredient_id)).all():
        exists=db.scalar(select(RecipeIngredient).where(RecipeIngredient.recipe_id==link.recipe_id,RecipeIngredient.ingredient_id==replacement.id))
        if exists:
            db.delete(link)
        else:
            link.ingredient_id=replacement.id
    ingredient.active=False
    db.commit()
    return RedirectResponse('/ingredients',303)

@app.post("/ingredients/{ingredient_id}/delete-with-recipes")
def delete_ingredient_and_disable_recipes(ingredient_id:int, db:Session=Depends(get_db)):
    ingredient=db.get(Ingredient,ingredient_id)
    if not ingredient: raise HTTPException(404)
    recipe_ids=[link.recipe_id for link in db.scalars(select(RecipeIngredient).where(RecipeIngredient.ingredient_id==ingredient_id)).all()]
    for recipe in db.scalars(select(Recipe).where(Recipe.id.in_(recipe_ids))).all():
        recipe.active=False
    ingredient.active=False
    db.commit()
    return RedirectResponse('/ingredients',303)

@app.get("/seasonings", response_class=HTMLResponse)
def seasonings_page(request:Request, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True).order_by(Branch.name)).all()
    selected_branch_id=int(request.query_params.get("branch_id") or 0) or None
    suppliers=filter_available(db.scalars(select(Supplier).where(Supplier.active==True).order_by(Supplier.name)).all(),"supplier",selected_branch_id,branches,db)
    items=filter_available(db.scalars(select(Seasoning).where(Seasoning.active==True).order_by(Seasoning.name)).all(),"seasoning",selected_branch_id,branches,db)
    supplier_map={s.id:s for s in suppliers}
    branch_map={b.id:b for b in branches}
    seasoning_scope_map=scope_map_for(items,"seasoning",branches,db)
    recipe_map={r.id:r for r in db.scalars(select(Recipe).where(Recipe.active==True)).all()}
    seasoning_recipes={}
    for link in db.scalars(select(RecipeSeasoning)).all():
        recipe=recipe_map.get(link.recipe_id)
        if recipe:
            seasoning_recipes.setdefault(link.seasoning_id,[]).append(recipe)
    return templates.TemplateResponse("seasonings.html",{"request":request,"branches":branches,"suppliers":suppliers,"items":items,"supplier_map":supplier_map,"branch_map":branch_map,"scope_map":seasoning_scope_map,"seasoning_recipes":seasoning_recipes,"selected_branch_id":selected_branch_id})

@app.post("/seasonings")
def create_seasoning(name:str=Form(...), supplier_id:int=Form(...), branch_id:str=Form(""), manufacturer:str=Form(""), product_name:str=Form(""), unit:str=Form(""), certification:str=Form(""), certification_no:str=Form(""), db:Session=Depends(get_db)):
    branch_value=optional_branch_id(branch_id)
    db.add(Seasoning(name=name,supplier_id=supplier_id,branch_id=branch_value,manufacturer=manufacturer or None,product_name=product_name or None,unit=unit or None,certification=certification or None,certification_no=certification_no or None))
    db.commit()
    return redirect_with_branch('/seasonings',branch_value)

@app.post("/seasonings/{seasoning_id}/scope")
async def update_seasoning_scope(seasoning_id:int, request:Request, db:Session=Depends(get_db)):
    seasoning=db.get(Seasoning,seasoning_id)
    if not seasoning: raise HTTPException(404)
    form=await request.form()
    branch_ids=[int(x) for x in form.getlist("branch_ids") if str(x).isdigit()]
    set_item_scopes(seasoning,"seasoning",branch_ids,db)
    db.commit()
    return redirect_with_branch('/seasonings',branch_ids[0] if branch_ids else None)

@app.post("/seasonings/{seasoning_id}/update")
def update_seasoning(seasoning_id:int, name:str=Form(...), supplier_id:int=Form(...), branch_id:str=Form(""), manufacturer:str=Form(""), product_name:str=Form(""), unit:str=Form(""), certification:str=Form(""), certification_no:str=Form(""), db:Session=Depends(get_db)):
    seasoning=db.get(Seasoning,seasoning_id)
    if not seasoning: raise HTTPException(404)
    branch_value=optional_branch_id(branch_id)
    seasoning.name=name.strip()
    seasoning.supplier_id=supplier_id
    seasoning.branch_id=branch_value
    seasoning.manufacturer=manufacturer.strip() or None
    seasoning.product_name=product_name.strip() or None
    seasoning.unit=unit.strip() or None
    seasoning.certification=certification.strip() or None
    seasoning.certification_no=certification_no.strip() or None
    db.commit()
    return redirect_with_branch('/seasonings',branch_value)

@app.post("/seasonings/{seasoning_id}/delete")
def delete_seasoning(seasoning_id:int, db:Session=Depends(get_db)):
    seasoning=db.get(Seasoning,seasoning_id)
    if not seasoning: raise HTTPException(404)
    for link in db.scalars(select(RecipeSeasoning).where(RecipeSeasoning.seasoning_id==seasoning_id)).all():
        db.delete(link)
    seasoning.active=False
    db.commit()
    return RedirectResponse('/seasonings',303)

@app.get("/menu", response_class=HTMLResponse)
def menu_page(request:Request, branch_id:int|None=None, db:Session=Depends(get_db)):
    branches=db.scalars(select(Branch).where(Branch.active==True)).all()
    recipes_all=db.scalars(select(Recipe).where(Recipe.active==True).order_by(Recipe.name)).all()
    recipes=filter_available(recipes_all,"recipe",branch_id,branches,db)
    items=[]
    if branch_id:
        items=db.scalars(select(DailyMenu).where(DailyMenu.branch_id==branch_id).order_by(DailyMenu.service_date.desc()).limit(100)).all()
    recipe_map={r.id:r for r in recipes}
    return templates.TemplateResponse("menu.html",{"request":request,"branches":branches,"recipes":recipes,"items":items,"recipe_map":recipe_map,"branch_id":branch_id})

@app.post("/menu/add-range")
def add_menu_range(branch_id:int=Form(...), recipe_id:int=Form(...), start_date:date=Form(...), end_date:date=Form(...), weekdays:str=Form("0,1,2,3,4"), servings:int|None=Form(None), db:Session=Depends(get_db)):
    allowed={int(x) for x in weekdays.split(',') if x!=''}
    d=start_date
    while d<=end_date:
        if d.weekday() in allowed:
            exists=db.scalar(select(DailyMenu).where(DailyMenu.branch_id==branch_id,DailyMenu.service_date==d,DailyMenu.recipe_id==recipe_id))
            if not exists:
                db.add(DailyMenu(branch_id=branch_id,service_date=d,recipe_id=recipe_id,servings=servings,upload_state='pending'))
        d+=timedelta(days=1)
    db.commit()
    return RedirectResponse(f'/menu?branch_id={branch_id}',303)

@app.post("/api/line/webhook")
async def line_webhook(request:Request):
    body=await request.body()
    config=LineConfig.from_env()
    signature=request.headers.get("X-Line-Signature")
    if not verify_line_signature(body, signature, config.channel_secret):
        raise HTTPException(status_code=403, detail="invalid LINE signature")
    payload=await request.json()
    sent=0
    errors=[]
    db=next(get_db())
    try:
        for event in payload.get("events",[]):
            text=handle_line_event(event, db, config)
            if text:
                try:
                    if reply_text(event.get("replyToken"), text, config):
                        sent+=1
                except Exception as exc:
                    errors.append(str(exc))
    finally:
        db.close()
    return JSONResponse({"ok":True,"received":len(payload.get("events",[])),"replied":sent,"errors":errors})

@app.get("/health")
def health(): return {"ok":True}

@app.post("/tasks/upload-reminders")
def run_upload_reminders_task(request:Request, db:Session=Depends(get_db)):
    result=send_due_upload_reminders(db)
    return {
        "ok": not result.errors,
        "target_date": result.target_date.isoformat(),
        "skipped": result.skipped,
        "sent": result.sent,
        "messages": result.messages,
        "errors": result.errors,
    }
