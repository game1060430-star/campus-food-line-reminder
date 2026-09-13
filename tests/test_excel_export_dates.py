from datetime import date

from openpyxl import Workbook, load_workbook

from app.models import Supplier
from app.services.excel_export import TEMPLATES, purchase_date_for, required_column_indexes, write_rows


def test_fixed_delivery_day_uses_previous_delivery_before_service_date():
    supplier = Supplier(name="喬富", delivery_weekdays="0,2")

    assert purchase_date_for(date(2026, 9, 14), supplier) == date(2026, 9, 9)
    assert purchase_date_for(date(2026, 9, 15), supplier) == date(2026, 9, 14)


def test_supplier_without_delivery_days_uses_previous_workday():
    supplier = Supplier(name="一般供應商")

    assert purchase_date_for(date(2026, 9, 14), supplier) == date(2026, 9, 11)


def test_write_rows_preserves_official_template_sheets_and_headers(tmp_path):
    path = tmp_path / "official.xlsx"
    path.write_bytes((TEMPLATES / "PreMenuExcelExample.xlsx").read_bytes())
    original = load_workbook(path, read_only=True)
    original_sheets = original.sheetnames
    original_headers = [original[original_sheets[0]].cell(1, col).value for col in range(1, 9)]
    original.close()

    write_rows(path, "菜色清單", [["A校", "午餐", "A餐廳", date(2026, 9, 14), "不應寫入", "蘑菇麵", "油麵", 550]], {4})

    result = load_workbook(path, read_only=True)
    assert result.sheetnames == original_sheets
    assert [result["菜色清單"].cell(1, col).value for col in range(1, 9)] == original_headers
    result.close()


def test_write_rows_leaves_unstarred_official_columns_blank(tmp_path):
    path = tmp_path / "official.xlsx"
    path.write_bytes((TEMPLATES / "PreMenuExcelExample.xlsx").read_bytes())

    write_rows(path, "菜色清單", [["A校", "午餐", "A餐廳", date(2026, 9, 14), "不應寫入", "蘑菇麵", "油麵", 550]], {4})

    result = load_workbook(path)
    sheet = result["菜色清單"]
    assert sheet.cell(2, 1).value == "A校"
    assert sheet.cell(2, 5).value is None
    assert sheet.cell(2, 6).value == "蘑菇麵"
    result.close()


def test_write_rows_can_include_menu_ingredient_composition(tmp_path):
    path = tmp_path / "official.xlsx"
    path.write_bytes((TEMPLATES / "PreMenuExcelExample.xlsx").read_bytes())

    write_rows(
        path,
        "菜色清單",
        [["A校", "午餐", "A餐廳", date(2026, 9, 14), "不應寫入", "蘑菇麵", "油麵、蘑菇醬", 550]],
        {4},
        {7},
    )

    result = load_workbook(path)
    sheet = result["菜色清單"]
    assert sheet.cell(2, 5).value is None
    assert sheet.cell(2, 7).value == "油麵、蘑菇醬"
    result.close()


def test_required_columns_fallback_to_all_columns_when_template_has_no_stars():
    wb = Workbook()
    ws = wb.active
    ws.append(["供應商名稱", "負責人"])

    assert required_column_indexes(ws) == {1, 2}
